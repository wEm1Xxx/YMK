from Models.Files_Practice import *
import io
from docx import Document


class TemplatesController:
    @staticmethod
    def get_all():
        return Files_Practice.select()

    @staticmethod
    def get(id):
        return Files_Practice.get_by_id(id)

    @staticmethod
    def debug_template(template_id):
        template = Files_Practice.get_by_id(template_id)
        doc = Document(io.BytesIO(template.file_data))
        placeholders = []
        for paragraph in doc.paragraphs:
            if '{' in paragraph.text and '}' in paragraph.text:
                placeholders.append(paragraph.text.strip())
        return placeholders